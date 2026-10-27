import sys
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
d = Document(sys.argv[1])
for t in d.tables:
    ncol = len(t.columns)
    size = Pt(8) if ncol >= 9 else Pt(9) if ncol >= 6 else Pt(10)
    t.autofit = True
    tblPr = t._tbl.tblPr
    for el in tblPr.findall(qn('w:tblW')):
        tblPr.remove(el)
    w = OxmlElement('w:tblW'); w.set(qn('w:w'), '5000'); w.set(qn('w:type'), 'pct'); tblPr.append(w)
    # bordes finos horizontales
    b = OxmlElement('w:tblBorders')
    for side in ('top', 'bottom', 'insideH'):
        e = OxmlElement(f'w:{side}'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'A6A6A6'); b.append(e)
    tblPr.append(b)
    for i, row in enumerate(t.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
                for r in p.runs:
                    r.font.size = size
                    if i == 0: r.bold = True
d.save(sys.argv[2])
